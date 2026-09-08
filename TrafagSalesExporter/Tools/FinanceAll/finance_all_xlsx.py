# /// script
# requires-python = ">=3.12,<3.13"
# dependencies = ["openpyxl>=3.1"]
# ///
"""Erzeugt Finance_All: Hauptbuch-Buchungszeilen aller Gesellschaften in einer Mappe.

Aufruf:
    uv run --python 3.12 Tools/FinanceAll/finance_all_xlsx.py                # alles
    uv run --python 3.12 Tools/FinanceAll/finance_all_xlsx.py --seit 2026-08-01
    uv run --python 3.12 Tools/FinanceAll/finance_all_xlsx.py --tage 30      # Schnelltest
    uv run --python 3.12 Tools/FinanceAll/finance_all_xlsx.py --lokal        # DB vorher kopieren

Bei grossen Auswertungen zuerst --lokal verwenden: Lesen ueber SMB war der Engpass.
Gemessen am 08.09.2026: 30 Tage 49 Sekunden, Vollauszug 4 min 10 s mit lokaler Kopie.
--tage/--seit begrenzen nur das Detailblatt; Summen und Feldstatus bleiben vollstaendig.

Das Layout der Detailzeilen folgt Andreas' Skizze `db1_fin.xlsx`: `entity` zuerst, damit
nach Gesellschaft sortiert und gefiltert werden kann. `due date` liest das importierte
Faelligkeitsdatum der Buchungszeile. Ohne Konzernkontenplan bleibt `Konzernkonto` leer.
Das Blatt Feldstatus zeigt die tatsaechliche Belegung je Gesellschaft.
"""
from __future__ import annotations

import argparse
import shutil
import sqlite3
import tempfile
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

PRODUKTIV_DB = r"\\trch-webapp-bidashboard.trafagch.local\BiDashboard$\trafag_exporter.db"
WURZEL = Path(__file__).resolve().parents[2]

KOPF_FARBE = PatternFill("solid", fgColor="1F4E79")
WARN_FARBE = PatternFill("solid", fgColor="FFF2CC")


def argumente():
    p = argparse.ArgumentParser(description="Finance_All erzeugen")
    p.add_argument("--seit", help="Nur Buchungen ab diesem Datum, Format JJJJ-MM-TT")
    p.add_argument("--tage", type=int, help="Nur die letzten N Tage; ueberschreibt --seit")
    p.add_argument("--db", default=PRODUKTIV_DB, help="Pfad zur Datenbank")
    p.add_argument("--lokal", action="store_true",
                   help="Datenbank vor dem Lesen lokal kopieren; ueber das Netz ist das oft schneller")
    p.add_argument("--ziel", help="Zieldatei; Vorgabe ist Finance_All_<heute>.xlsx in der Wurzel")
    return p.parse_args()


def kopfzeile(ws, spalten, breiten):
    zellen = []
    for s in spalten:
        z = WriteOnlyCell(ws, value=s)
        z.font = Font(bold=True, color="FFFFFF")
        z.fill = KOPF_FARBE
        z.alignment = Alignment(vertical="center", wrap_text=True)
        zellen.append(z)
    ws.append(zellen)
    for i, b in enumerate(breiten, start=1):
        ws.column_dimensions[get_column_letter(i)].width = b
    ws.freeze_panes = "A2"


def hinweiszeile(ws, text, art=None):
    z = WriteOnlyCell(ws, value=text)
    z.alignment = Alignment(wrap_text=True, vertical="top")
    if art == "titel":
        z.font = Font(bold=True, size=16, color="1F4E79")
    elif art == "klein":
        z.font = Font(size=9, italic=True, color="595959")
    elif art == "h":
        z.font = Font(bold=True, size=11)
    elif art == "warn":
        z.fill = WARN_FARBE
    ws.append([z])


def main() -> None:
    a = argumente()

    seit = None
    if a.tage:
        seit = (date.today() - timedelta(days=a.tage)).isoformat()
    elif a.seit:
        seit = a.seit

    quelle = a.db
    tempdir = None
    if a.lokal:
        tempdir = tempfile.mkdtemp(prefix="finance_all_")
        kopie = Path(tempdir) / "trafag_exporter.db"
        print(f"Datenbank wird lokal kopiert nach {kopie} ...")
        shutil.copy2(a.db, kopie)
        quelle = str(kopie)

    ziel = Path(a.ziel) if a.ziel else WURZEL / f"Finance_All_{date.today():%Y-%m-%d}.xlsx"

    con = sqlite3.connect(f"file:{quelle}?mode=ro", uri=True)
    con.row_factory = sqlite3.Row
    journal_spalten = {r["name"] for r in con.execute("PRAGMA table_info(FinancialJournalEntries)")}
    # Alte Produktivstaende bleiben lesbar, bis Schema und Journalimport erneuert sind.
    due_sql = "DueDate" if "DueDate" in journal_spalten else "NULL"

    where = ""
    parameter: tuple = ()
    if seit:
        where = " WHERE PostingDate >= ?"
        parameter = (seit,)

    wb = Workbook(write_only=True)

    # --- Lesehinweis ---------------------------------------------------------
    ws = wb.create_sheet("Lesehinweis")
    ws.column_dimensions["A"].width = 118
    hinweiszeile(ws, "Finance_All - Hauptbuch aller Gesellschaften", "titel")
    hinweiszeile(ws, f"Stand {datetime.now():%d.%m.%Y %H:%M}, gelesen aus der zentralen Datenbank", "klein")
    hinweiszeile(ws, "")
    hinweiszeile(ws, "Was diese Mappe ist", "h")
    hinweiszeile(ws,
                 "Die Finanzsicht: eine Zeile je Buchungszeile aus dem Hauptbuch. Das Gegenstueck zu "
                 "Sales_All, das die Vertriebssicht mit der Rechnungsposition als kleinster Einheit "
                 "zeigt. Beides gehoert nicht in eine Datei: eine Rechnung erzeugt mehrere "
                 "Buchungszeilen, und Zahlungen, Lagerbuchungen, Abschreibungen und manuelle "
                 "Buchungen haben ueberhaupt keine Verkaufszeile.")
    hinweiszeile(ws, "")
    if seit:
        hinweiszeile(ws, "ACHTUNG: gekuerzter Auszug", "h")
        hinweiszeile(ws,
                     f"Das Blatt 'Journal Detail' enthaelt nur Buchungen ab {seit}. Die Blaetter "
                     "'Journal Summary', 'Konten' und 'Datenstatus' sind vollstaendig. Fuer die "
                     "vollstaendige Mappe das Skript ohne --seit und ohne --tage aufrufen.", "warn")
        hinweiszeile(ws, "")
    hinweiszeile(ws, "Blaetter", "h")
    hinweiszeile(ws, "Journal Detail - Buchungszeilen. Spalte entity zuerst, damit nach Gesellschaft "
                     "sortiert und gefiltert werden kann.")
    hinweiszeile(ws, "Journal Summary - Summen je Gesellschaft, Konto und Periode.")
    hinweiszeile(ws, "Konten - jedes Konto je Gesellschaft einmal. Das ist die Arbeitsliste fuer das "
                     "Konzernkonto-Mapping: Spalte Konzernkonto ausfuellen und zurueckgeben.")
    hinweiszeile(ws, "Datenstatus - welche Gesellschaft wie weit geladen ist.")
    hinweiszeile(ws, "Feldstatus - Belegung je Gesellschaft im gesamten geladenen Journalbestand.")
    hinweiszeile(ws, "")
    hinweiszeile(ws, "Was fehlt, und zwar sichtbar", "h")
    hinweiszeile(ws, "Die Schweiz und Oesterreich fehlen. Ihr Hauptbuch kommt aus SAP; das dortige "
                     "EntitySet heisst FinanzdataSchweizOeSet, der Leser sucht aber weiterhin nach "
                     "FinanzJournalSet. Siehe ISS-006.", "warn")
    hinweiszeile(ws, "due date zeigt das Faelligkeitsdatum der Buchungszeile, kein Zahlungsdatum. "
                     "Nach Erweiterung des Imports muessen die Gesellschaften erneut geladen werden; "
                     "die Belegung steht im Blatt Feldstatus.")
    hinweiszeile(ws, "Der Konzernkontenplan mit Zuordnungen je Gesellschaft liegt noch nicht vor; "
                     "Konzernkonto bleibt deshalb leer. Kostenstelle und dimension 2 werden aus "
                     "den Buchungszeilen uebernommen; leere Quellfelder bleiben leer.", "warn")
    hinweiszeile(ws, "")
    hinweiszeile(ws, "Betraege und Waehrungen", "h")
    hinweiszeile(ws, "Die Spalte amount ist der Betrag mit Vorzeichen in der Landeswaehrung: Soll "
                     "positiv, Haben negativ. Es gibt bewusst keine Summe ueber Gesellschaften "
                     "hinweg, weil die Waehrungen unterschiedlich sind.")
    hinweiszeile(ws, "Das Kennzeichen db/cr folgt der Skizze vom 08.09.2026: 40 fuer Soll, 50 fuer Haben.")

    # --- Journal Detail ------------------------------------------------------
    ws = wb.create_sheet("Journal Detail")
    kopfzeile(ws,
              ["entity", "db/cr", "accnt", "acct desrc", "amount", "currency", "entry text",
               "posting date", "entry ID", "line ID", "due date", "Konzernkonto",
               "fiscal year", "period", "cost center", "dimension 2", "doc type",
               "source document", "manual", "reversal", "source system"],
              [10, 8, 14, 34, 16, 10, 42, 13, 16, 9, 12, 14, 11, 8, 14, 14, 10, 18, 9, 9, 13])

    anzahl = 0
    for r in con.execute(
            "SELECT Tsc, CAST(DebitAmount AS REAL) AS Soll, CAST(CreditAmount AS REAL) AS Haben,"
            " AccountCode, AccountName, CAST(SignedAmountLocal AS REAL) AS Betrag,"
            " LocalCurrency, LineMemo, PostingDate, JournalEntryId, JournalEntryLineId,"
            " FiscalYear, FiscalPeriod, CostCenter, Dimension2, TransactionType,"
            " SourceDocumentNumber, IsManual, IsReversal, SourceSystem, " + due_sql + " AS DueDate"
            " FROM FinancialJournalEntries" + where +
            " ORDER BY Tsc, PostingDate, JournalEntryId, JournalEntryLineId", parameter):
        ws.append([
            r["Tsc"],
            40 if (r["Soll"] or 0) > 0 else 50,
            r["AccountCode"], r["AccountName"], r["Betrag"], r["LocalCurrency"], r["LineMemo"],
            (r["PostingDate"] or "")[:10],
            r["JournalEntryId"], r["JournalEntryLineId"],
            (r["DueDate"] or "")[:10], "",
            r["FiscalYear"], r["FiscalPeriod"], r["CostCenter"], r["Dimension2"],
            r["TransactionType"], r["SourceDocumentNumber"],
            "ja" if r["IsManual"] else "nein",
            "ja" if r["IsReversal"] else "nein",
            r["SourceSystem"],
        ])
        anzahl += 1

    # --- Journal Summary -----------------------------------------------------
    ws = wb.create_sheet("Journal Summary")
    kopfzeile(ws, ["entity", "accnt", "acct desrc", "fiscal year", "period", "currency",
                   "Soll", "Haben", "Saldo", "Zeilen", "Konzernkonto"],
              [10, 14, 34, 11, 8, 10, 16, 16, 16, 10, 14])
    for r in con.execute(
            "SELECT Tsc, AccountCode, MAX(AccountName) AS Name, FiscalYear, FiscalPeriod,"
            " MAX(LocalCurrency) AS Waehrung, SUM(CAST(DebitAmount AS REAL)) AS Soll,"
            " SUM(CAST(CreditAmount AS REAL)) AS Haben,"
            " SUM(CAST(SignedAmountLocal AS REAL)) AS Saldo, COUNT(*) AS Zeilen"
            " FROM FinancialJournalEntries"
            " GROUP BY Tsc, AccountCode, FiscalYear, FiscalPeriod"
            " ORDER BY Tsc, AccountCode, FiscalYear, FiscalPeriod"):
        ws.append([r["Tsc"], r["AccountCode"], r["Name"], r["FiscalYear"], r["FiscalPeriod"],
                   r["Waehrung"], r["Soll"], r["Haben"], r["Saldo"], r["Zeilen"], ""])

    # --- Konten --------------------------------------------------------------
    ws = wb.create_sheet("Konten")
    kopfzeile(ws, ["entity", "accnt", "acct desrc", "currency", "Zeilen", "Saldo gesamt",
                   "Konzernkonto", "Konzernkontobezeichnung"],
              [10, 14, 40, 10, 10, 18, 14, 34])
    for r in con.execute(
            "SELECT Tsc, AccountCode, MAX(AccountName) AS Name, MAX(LocalCurrency) AS Waehrung,"
            " COUNT(*) AS Zeilen, SUM(CAST(SignedAmountLocal AS REAL)) AS Saldo"
            " FROM FinancialJournalEntries GROUP BY Tsc, AccountCode ORDER BY Tsc, AccountCode"):
        ws.append([r["Tsc"], r["AccountCode"], r["Name"], r["Waehrung"], r["Zeilen"], r["Saldo"], "", ""])

    # --- Datenstatus ---------------------------------------------------------
    ws = wb.create_sheet("Datenstatus")
    kopfzeile(ws, ["entity", "Quellsystem", "Zeilen", "Buchungsdatum von", "bis",
                   "zuletzt geladen", "Hinweis"],
              [10, 14, 12, 18, 14, 22, 62])
    vorhanden = set()
    for r in con.execute(
            "SELECT Tsc, MAX(SourceSystem) AS Quelle, COUNT(*) AS Zeilen, MIN(PostingDate) AS Von,"
            " MAX(PostingDate) AS Bis, MAX(StoredAtUtc) AS Geladen"
            " FROM FinancialJournalEntries GROUP BY Tsc ORDER BY Tsc"):
        vorhanden.add(r["Tsc"])
        ws.append([r["Tsc"], r["Quelle"], r["Zeilen"], (r["Von"] or "")[:10], (r["Bis"] or "")[:10],
                   (r["Geladen"] or "")[:19], ""])
    if "ZSCHWEIZ" not in vorhanden:
        ws.append(["ZSCHWEIZ", "SAP", 0, "", "", "",
                   "Fehlt: der Leser sucht FinanzJournalSet, in P76 heisst das EntitySet "
                   "FinanzdataSchweizOeSet (ISS-006)"])

    # --- Feldstatus (vollstaendiger Bestand, auch beim gekuerzten Detailblatt) --
    ws = wb.create_sheet("Feldstatus")
    kopfzeile(ws, ["entity", "Feld", "Zeilen gesamt", "Gefuellt", "Leer", "Hinweis"],
              [10, 18, 16, 14, 14, 80])
    felder = [
        ("fiscal year", "FiscalYear", "Geschaeftsjahr laut Import."),
        ("period", "FiscalPeriod", "Periode laut Import."),
        ("cost center", "CostCenter", "B1: Verteilungsregel der Dimension 1; ohne Quellwert leer."),
        ("dimension 2", "Dimension2", "B1: Verteilungsregel der Dimension 2; ohne Quellwert leer."),
        ("due date", due_sql, "Faelligkeitsdatum; fehlende Werte ggf. durch erneutes Laden nachziehen."),
        ("Konzernkonto", "NULL", "Wartet auf Konzernkontenplan und Zuordnung je Gesellschaft."),
    ]
    zaehler = ", ".join(
        f"SUM(CASE WHEN CAST({spalte} AS REAL) > 0 THEN 1 ELSE 0 END) AS f{i}"
        if feld in ("fiscal year", "period") else
        f"SUM(CASE WHEN NULLIF(TRIM({spalte}), '') IS NOT NULL THEN 1 ELSE 0 END) AS f{i}"
        for i, (feld, spalte, _) in enumerate(felder))
    for r in con.execute("SELECT Tsc, COUNT(*) AS Gesamt, " + zaehler +
                         " FROM FinancialJournalEntries GROUP BY Tsc ORDER BY Tsc"):
        for i, (feld, _, hinweis) in enumerate(felder):
            ws.append([r["Tsc"], feld, r["Gesamt"], r[f"f{i}"],
                       r["Gesamt"] - r[f"f{i}"], hinweis])

    wb.save(ziel)
    con.close()
    if tempdir:
        shutil.rmtree(tempdir, ignore_errors=True)

    umfang = f"ab {seit}" if seit else "vollstaendig"
    print(f"geschrieben: {ziel.name} | {anzahl} Buchungszeilen ({umfang}) "
          f"| {ziel.stat().st_size / 1024 / 1024:.1f} MB")


if __name__ == "__main__":
    main()
