#!/usr/bin/env python3
"""Baut Board (HTML), PDF und Excel fuer den Liefertreue-Monatsreport Trafag AG (VKORG 1100).
Aufruf: python3 build_liefertreue.py lt.json board_template.html <outdir>
lt.json = JSON.stringify(window.__lt) aus extract_liefertreue.js
Erzeugt: <outdir>/board.html (Artifact-Seite ohne doctype), Liefertreue_Trafag_1100_JJJJ-MM.pdf, Liefertreue_1100_verspaetet_JJJJ-MM.xlsx
"""
import json, sys, os, datetime

MN = ["Jan", "Feb", "Mär", "Apr", "Mai", "Jun", "Jul", "Aug", "Sep", "Okt", "Nov", "Dez"]
MF = ["Januar", "Februar", "März", "April", "Mai", "Juni", "Juli", "August", "September", "Oktober", "November", "Dezember"]


def main(src, tpl, out):
    d = json.load(open(src, encoding='utf-8'))
    if d.get('err'):
        sys.exit('Abzug fehlerhaft: ' + d['err'])
    cut, year = d['cut'], d['year']
    lm = int(cut[4:6])
    stand = datetime.date.fromisoformat(d['stand']).strftime('%d.%m.%Y')
    period = f"Jan – {MN[lm-1]}" if lm > 1 else "Jan"
    otd = [r.split(';') for r in d['otd']]
    otp = [r.split(';') for r in d['otp']]
    col = lambda a, i: '[' + ','.join(x[i] for x in a) + ']'
    months = [MN[int(x[0][4:6]) - 1] for x in otd]
    D = ("const D={\n  m:" + json.dumps(months, ensure_ascii=False) +
         f",\n  items:{col(otd,1)}, strict:{col(otd,2)}, tol:{col(otd,3)}, early:{col(otd,7)},\n"
         f"  orders:{col(otd,4)}, otifS:{col(otd,5)}, otifT:{col(otd,6)},\n"
         f"  pItems:{col(otp,1)}, pS:{col(otp,2)}, pT:{col(otp,3)}\n}};")
    PG = 'const PG=' + json.dumps(d['pg'], ensure_ascii=False, separators=(',', ':')) + ';'
    nc, npn = d['names']['c'], d['names']['p']
    L = [x.split(';') for x in d['list']]
    cus = sorted({x[6] for x in L}); ci = {c: i for i, c in enumerate(cus)}
    prs = sorted({x[7] for x in L}); pi = {p: i for i, p in enumerate(prs)}
    toi = lambda v: '' if v == '' else int(v)
    rows = [[x[0], x[1], x[2], x[3], toi(x[4]), x[5], ci[x[6]], pi[x[7]], int(x[8]), x[9], x[10], x[11], toi(x[12]), int(x[13]), x[14]] for x in L]
    LATE = 'const LATE=' + json.dumps(dict(cus=[[c, nc.get(c, '')] for c in cus], prs=[[p, npn.get(p, '')] for p in prs], late=rows), ensure_ascii=False, separators=(',', ':')) + ';'
    s = open(tpl, encoding='utf-8').read()
    nxt = MF[lm % 12]
    for k, v in {'__PERIOD__': period, '__YEAR__': year, '__STAND__': stand, '__LASTM__': MN[lm-1], '__LASTMONTHNAME__': MF[lm-1], '__NEXTMONTHNAME__': nxt}.items():
        s = s.replace(k, v)
    full = s.replace('/*__D__*/', D).replace('/*__PG__*/', PG).replace('/*__LATE__*/', LATE)
    os.makedirs(out, exist_ok=True)
    open(os.path.join(out, 'board.html'), 'w', encoding='utf-8').write(full)
    tag = f"{year}-{cut[4:6]}"
    # PDF
    pdf = full.replace('<style>', '<style>@page{size:A4;margin:12mm} body{font-size:12px} .wrap{gap:14px;padding-block:0} .kpi .v{font-size:1.8rem} .kpis{grid-template-columns:repeat(4,1fr)!important} .filters,.more,.count{display:none!important} #tblL,#tblM,#tblG,#tblW{min-width:0;font-size:.7rem} section.card{break-inside:avoid} #liste,#wunsch{break-inside:auto} .rank button{padding:2px 6px}', 1)
    pdf = (pdf.replace('<h2>Liste der Abweichungen (OTD)</h2>', f'<h2>Die 25 grössten Verzüge in der Auswertung (OTD, {period})</h2>')
           .replace("const st={q:'',m:'',t:'L',s:'',i:'',x:'',b:-1,cu:'',pr:'',sort:'wds',dir:-1,lim:100};", "const st={q:'',m:'',t:'L',s:'',i:'',x:'in',b:-1,cu:'',pr:'',sort:'wds',dir:-1,lim:25};")
           .replace(' · Klick filtert die Liste', '')
           .replace("const gs={q:'',min:20,sort:'n',dir:-1};", "const gs={q:'',min:100,sort:'n',dir:-1};")
           .replace(' Klick auf eine Spaltenüberschrift sortiert.', ' Im PDF: Gruppen ab 100 Positionen.'))
    ph = os.path.join(out, 'pdf.html')
    open(ph, 'w', encoding='utf-8').write('<!doctype html><html lang="de"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width">' + pdf + '</html>')
    pdf_path = os.path.join(out, f'Liefertreue_Trafag_1100_{tag}.pdf')
    from playwright.sync_api import sync_playwright
    with sync_playwright() as p:
        exe = '/opt/pw-browsers/chromium' if os.path.exists('/opt/pw-browsers/chromium') else None
        b = p.chromium.launch(executable_path=exe) if exe else p.chromium.launch()
        pg = b.new_page()
        pg.goto('file://' + os.path.abspath(ph)); pg.wait_for_timeout(1500)
        pg.emulate_media(color_scheme='light')
        pg.pdf(path=pdf_path, format='A4', print_background=True, margin={'top': '12mm', 'bottom': '12mm', 'left': '10mm', 'right': '10mm'})
        b.close()
    # Excel
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook(); ws = wb.active; ws.title = 'Abweichungen OTD'
    ws.append(['Art', 'Auftrag', 'Position', 'Kunde Nr.', 'Kunde', 'Kundengruppe', 'In Auswertung', 'Ausschlussgrund', 'Material', 'Materialtext', 'Bestätigter Termin', 'Monat', 'Ist-WA (letzter)', 'Abweichung AT', 'Wunschliefertermin', 'Abweichung Wunsch AT', 'Status', 'Wert CHF', 'Hinweis'])
    stt = {'C': 'geliefert', 'B': 'teilgeliefert', 'A': 'offen'}; xr = {'K': 'Gesellschaft ausgeschlossen', 'Z': 'Zertifikat EN10204', '': ''}
    dt = lambda s: datetime.date(int(s[:4]), int(s[4:6]), int(s[6:])) if s and s > '1990' else None
    for x in L:
        c = dt(x[2])
        ws.append(['zu spät' if x[14] == 'L' else 'zu früh', int(x[0]), int(x[1]), x[6], nc.get(x[6], ''), 'Trafag-Gesellschaft' if x[9] == 'I' else 'Extern', 'nein' if x[10] else 'ja', xr[x[10]], x[7], npn.get(x[7], ''), c, c.strftime('%Y-%m') if c else '', dt(x[3]), int(x[4]) if x[4] != '' else None, dt(x[11]), int(x[12]) if x[12] != '' else None, stt.get(x[5], x[5]), int(x[8]), 'kein WA gefunden, SAP-Kennzahl' if x[13] == '1' else ''])
    for c in ws[1]: c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='1F5F8B')
    for i, w in enumerate([9, 10, 8, 10, 38, 18, 12, 24, 12, 38, 12, 9, 12, 12, 12, 14, 13, 11, 26], 1): ws.column_dimensions[get_column_letter(i)].width = w
    for cl in 'KMO':
        for c in ws[cl][1:]: c.number_format = 'DD.MM.YYYY'
    for c in ws['R'][1:]: c.number_format = '#,##0'
    ws.freeze_panes = 'A2'; ws.auto_filter.ref = ws.dimensions
    d2 = wb.create_sheet('Definition')
    for l in ['Eigene Berechnung: letzter Ist-Warenausgang aller Lieferungen der Position (C_DELIVITMPERFANLYTSQRY) gegen bestätigten Liefertermin bzw. Wunschtermin (C_SLSORDDELIVPERFANLYTSQRY). Arbeitstage Mo–Fr ohne Feiertage.',
              'OTD: vollständig geliefert und Ist-WA höchstens 4 AT nach bestätigtem Termin. OTIF: alle Positionen des Auftrags OTD. OTP: wie OTD, gegen Wunschliefertermin.',
              'Zu spät: ≥ 5 AT nach bestätigtem Termin oder offen/teilgeliefert. Zu früh: Ist-WA ≥ 5 AT vor bestätigtem Termin.',
              'Nicht in der Auswertung: Trafag Spol. s.r.o. (CZ), Trafag Controls India, GFS Germany, Zertifikate EN10204.',
              f'Zeitraum: bestätigter Termin {period} {year}, VKORG 1100, Stand {stand}.']:
        d2.append([l])
    d2.column_dimensions['A'].width = 150
    xl_path = os.path.join(out, f'Liefertreue_1100_verspaetet_{tag}.xlsx'); wb.save(xl_path)
    # Kurzfazit
    S = lambda a, i: sum(int(x[i]) for x in a)
    n, st_, tl, on, os_, ot = S(otd, 1), S(otd, 2), S(otd, 3), S(otd, 4), S(otd, 5), S(otd, 6)
    pn, ps, pt = S(otp, 1), S(otp, 2), S(otp, 3)
    last = otd[-1]; lastp = otp[-1]
    summary = {'period': f'{period} {year}', 'stand': stand,
               'OTD': [round(100*tl/n, 1), round(100*st_/n, 1), n], 'OTIF': [round(100*ot/on, 1), round(100*os_/on, 1), on], 'OTP': [round(100*pt/pn, 1), round(100*ps/pn, 1), pn],
               'last_month': MN[lm-1], 'OTD_last': round(100*int(last[3])/int(last[1]), 1), 'OTIF_last': round(100*int(last[6])/max(1, int(last[4])), 1), 'OTP_last': round(100*int(lastp[3])/int(lastp[1]), 1),
               'early5': S(otd, 7), 'late_in_eval': sum(1 for x in L if not x[10] and x[14] == 'L'), 'pdf': pdf_path, 'xlsx': xl_path}
    print(json.dumps(summary, ensure_ascii=False, indent=1))


if __name__ == '__main__':
    main(*sys.argv[1:4])
