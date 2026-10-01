# /// script
# requires-python = ">=3.12,<3.13"
# dependencies = ["openpyxl>=3.1"]
# ///
"""Schreibt `Wochen_Todo.tsv` und `Vorhaben_HERMES.tsv` in `Wochen_Todo.xlsx` fort.

Aufruf:
    uv run --python 3.12 projektmanagement/wochen_todo_xlsx.py

Die TSV-Dateien sind die fuehrenden Quellen, die Excel nur die lesbare Ausgabe. Das
Skript **oeffnet die vorhandene Datei und ersetzt nur die Blaetter, die es selbst
erzeugt**, statt eine neue Mappe zu schreiben. Sonst ginge das Blatt `Erläuterungen`
verloren, das am 2026-09-02 dazugekommen ist und nicht aus der TSV stammt.

Erzeugte Blaetter:
- `Wochen-Todo`: die Aufgabenliste, seit 2026-09-29 mit den Spalten `Vorhaben` und
  `HERMES-Stufe`.
- `HERMES Übersicht`: eine Zeile je Vorhaben fuer die Geschaeftsleitung, mit Ampel,
  Rollen, den drei Entscheidungspunkten, Kennzahlkacheln und zwei Diagrammen. Steht
  als erstes Blatt.
- `Vorlage Projektauftrag` und `Vorlage Abnahmeprotokoll`: je eine Seite zum Ausfuellen.

Alle uebrigen Blaetter bleiben unveraendert.

Zusaetzlich, seit 2026-10-01, eine EIGENE Datei `Status_kurz.xlsx` fuer den Chef (Wunsch Ingo:
separates Excel, schoen dargestellt). Quelle ist `Status_kurz.tsv`, wenige Zeilen je Bereich mit
Bereit / Laeuft / Wartet / Offen; die Zahlen oben werden aus `Wochen_Todo.tsv` gerechnet, damit
die vielen kleinen und fremdabhaengigen Punkte nicht wie 50 Baustellen aussehen. Die Datei wird
bei jedem Lauf neu geschrieben.
"""

from __future__ import annotations

import csv
import math
from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.series import DataPoint
from openpyxl.formatting.rule import DataBarRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ORDNER = Path(__file__).resolve().parent
TSV = ORDNER / "Wochen_Todo.tsv"
VORHABEN_TSV = ORDNER / "Vorhaben_HERMES.tsv"
KURZ_TSV = ORDNER / "Status_kurz.tsv"
KURZ_XLSX = ORDNER / "Status_kurz.xlsx"
XLSX = ORDNER / "Wochen_Todo.xlsx"

BLATT_TODO = "Wochen-Todo"
BLATT_HERMES = "HERMES Übersicht"
BLATT_AUFTRAG = "Vorlage Projektauftrag"
BLATT_ABNAHME = "Vorlage Abnahmeprotokoll"

BREITEN = {
    "Nr": 5, "Dringlichkeit": 14, "Bereich": 20, "ID": 8, "Thema": 30,
    "Wer": 10, "Prio": 8, "Status": 12, "Fällig": 12,
    "Nächster Schritt": 90, "Erledigt am": 13, "Vorhaben": 10, "HERMES-Stufe": 14,
}

FARBEN = {
    "Erledigt": "C6E0B4",
    "Läuft": "FFE699",
    "Offen": "F8CBAD",
}

# Status kurz: Hintergrund je Stand
KURZ_FARBEN = {"Bereit": "C6E0B4", "Läuft": "FFE699", "Wartet": "DDEBF7", "Offen": "F8CBAD"}
KURZ_REIHE = ("Bereit", "Läuft", "Wartet", "Offen")
KURZ_TEXT = {"Bereit": "bereit", "Läuft": "läuft", "Wartet": "wartet auf andere", "Offen": "offen"}

STUFE_FARBEN = {"Kleinauftrag": "E2EFDA", "Vorhaben": "DDEBF7", "Projekt": "E4DFEC"}

# Ampel: Symbolfarbe, Hintergrund
AMPEL = {
    "Grün": ("2E8B57", "E2F0D9"),
    "Gelb": ("C9A100", "FFF2CC"),
    "Rot": ("C00000", "FCE4E4"),
    "Neu": ("7F7F7F", "EDEDED"),
}
AMPEL_REIHE = ("Grün", "Gelb", "Rot", "Neu")

DUNKEL = "1F3864"
HELL = "D9E1F2"
GRAU = "7F7F7F"
SCHRIFT = "Segoe UI"
DUENN = Side(style="thin", color="BFBFBF")
RAHMEN = Border(left=DUENN, right=DUENN, top=DUENN, bottom=DUENN)


def lies_tsv(pfad: Path) -> list[list[str]]:
    with pfad.open(encoding="utf-8-sig", newline="") as f:
        return [z for z in csv.reader(f, delimiter="\t") if any(z)]


def als_dicts(zeilen: list[list[str]]) -> list[dict[str, str]]:
    kopf = zeilen[0]
    return [{k: (z[i] if i < len(z) else "") for i, k in enumerate(kopf)} for z in zeilen[1:]]


def neues_blatt(mappe, name: str, index: int | None = None):
    """Loescht ein selbst erzeugtes Blatt und legt es an derselben Stelle neu an."""
    if name in mappe.sheetnames:
        alt = mappe.sheetnames.index(name)
        del mappe[name]
        if index is None:
            index = alt
    return mappe.create_sheet(name, index)


def zelle(blatt, ref: str, wert, **stil):
    z = blatt[ref]
    z.value = wert
    for k, v in stil.items():
        setattr(z, k, v)
    return z


def schreibe_todo(mappe, zeilen: list[list[str]]) -> None:
    kopf = zeilen[0]
    blatt = mappe[BLATT_TODO]
    blatt.delete_rows(1, blatt.max_row)
    for zeile in zeilen:
        blatt.append(zeile)

    for spalte, name in enumerate(kopf, start=1):
        z = blatt.cell(row=1, column=spalte)
        z.font = Font(bold=True)
        z.alignment = Alignment(vertical="top")
        blatt.column_dimensions[get_column_letter(spalte)].width = BREITEN.get(name, 18)

    status_spalte = kopf.index("Status") + 1 if "Status" in kopf else None
    stufe_spalte = kopf.index("HERMES-Stufe") + 1 if "HERMES-Stufe" in kopf else None

    for nr in range(2, blatt.max_row + 1):
        for spalte in range(1, len(kopf) + 1):
            blatt.cell(row=nr, column=spalte).alignment = Alignment(
                vertical="top", wrap_text=True)
        for spalte, farben in ((status_spalte, FARBEN), (stufe_spalte, STUFE_FARBEN)):
            if not spalte:
                continue
            wert = blatt.cell(row=nr, column=spalte).value
            farbe = farben.get(str(wert).strip()) if wert else None
            if farbe:
                blatt.cell(row=nr, column=spalte).fill = PatternFill("solid", fgColor=farbe)

    blatt.freeze_panes = "A2"
    blatt.auto_filter.ref = blatt.dimensions


def lies_datum(wert: str) -> date | None:
    """Erledigt am / Seit kommen als JJJJ-MM-TT (mit oder ohne Uhrzeit) oder TT.MM.JJJJ."""
    wert = (wert or "").strip()
    try:
        if len(wert) >= 10 and wert[4] == "-":
            return date.fromisoformat(wert[:10])
        if len(wert) >= 10 and wert[2] == ".":
            return date(int(wert[6:10]), int(wert[3:5]), int(wert[0:2]))
    except ValueError:
        return None
    return None


def schreibe_status_kurz(punkte: list[dict[str, str]], aufgaben: list[dict[str, str]]) -> None:
    """Eigene Datei fuer den Chef: je Bereich wenige Zeilen, Details stehen in Wochen_Todo.xlsx."""
    mappe = Workbook()
    blatt = mappe.active
    blatt.title = "Status"
    heute = date.today()
    for spalte, breite in zip("ABCDE", (2, 58, 13, 34, 12)):
        blatt.column_dimensions[spalte].width = breite

    # Kopfband
    for spalte in range(1, 6):
        for zeile in (1, 2):
            blatt.cell(row=zeile, column=spalte).fill = PatternFill("solid", fgColor=DUNKEL)
    zelle(blatt, "B1", "Informatik / Analytics  ·  Status kurz",
          font=Font(name=SCHRIFT, size=18, bold=True, color="FFFFFF"),
          alignment=Alignment(vertical="center"))
    zelle(blatt, "B2", f"Stand {heute:%d.%m.%Y}  ·  Ingo Kohler",
          font=Font(name=SCHRIFT, size=10, color="D9E1F2"), alignment=Alignment(vertical="top"))
    blatt.row_dimensions[1].height = 34
    blatt.row_dimensions[2].height = 20

    # Kacheln: Zahl gross, Text klein
    zeile = 4
    for i, stand in enumerate(KURZ_REIHE):
        anzahl = sum(1 for p in punkte if p.get("Stand") == stand)
        spalte = (2, 3, 4, 5)[i]
        oben = blatt.cell(row=zeile, column=spalte, value=anzahl)
        oben.font = Font(name=SCHRIFT, size=20, bold=True, color=DUNKEL)
        oben.alignment = Alignment(horizontal="center", vertical="center")
        unten = blatt.cell(row=zeile + 1, column=spalte, value=KURZ_TEXT[stand])
        unten.font = Font(name=SCHRIFT, size=9, color=DUNKEL)
        unten.alignment = Alignment(horizontal="center")
        for z in (oben, unten):
            z.fill = PatternFill("solid", fgColor=KURZ_FARBEN[stand])
    blatt.row_dimensions[zeile].height = 32

    offen = [a for a in aufgaben if a.get("Status", "").strip() != "Erledigt"]
    wartet = [a for a in offen
              if a.get("Status", "").startswith("Wartet") or "Ingo" not in a.get("Wer", "")]
    klein = [a for a in offen if a.get("Dringlichkeit", "") == "Nur terminieren"]
    erledigt = [a for a in aufgaben if a.get("Status", "").strip() == "Erledigt"
                and (d := lies_datum(a.get("Erledigt am", ""))) and (heute - d).days <= 14]
    zelle(blatt, "B7",
          f"Hinter dieser Seite stehen {len(offen)} offene Detailpunkte (Wochen_Todo.xlsx). "
          f"Davon warten {len(wartet)} auf andere, {len(klein)} sind Kleinigkeiten zum Terminieren. "
          f"In den letzten 14 Tagen erledigt: {len(erledigt)}.",
          font=Font(name=SCHRIFT, size=10, italic=True, color=GRAU),
          alignment=Alignment(wrap_text=True, vertical="top"))
    blatt.merge_cells("B7:E7")
    blatt.row_dimensions[7].height = 30

    zeile = 9
    for spalte, name in zip(range(2, 6), ("Punkt", "Stand", "Wartet auf", "Seit")):
        z = blatt.cell(row=zeile, column=spalte, value=name)
        z.font = Font(name=SCHRIFT, bold=True, color="FFFFFF")
        z.fill = PatternFill("solid", fgColor=DUNKEL)
        z.alignment = Alignment(vertical="center")
    blatt.row_dimensions[zeile].height = 20
    zeile += 1

    bereiche: list[str] = []
    for p in punkte:
        if p.get("Bereich") not in bereiche:
            bereiche.append(p.get("Bereich", ""))
    linie = Border(bottom=Side(style="thin", color="D9D9D9"))
    for bereich in bereiche:
        z = blatt.cell(row=zeile, column=2, value=bereich)
        z.font = Font(name=SCHRIFT, size=12, bold=True, color=DUNKEL)
        for spalte in range(2, 6):
            blatt.cell(row=zeile, column=spalte).fill = PatternFill("solid", fgColor=HELL)
        blatt.row_dimensions[zeile].height = 22
        zeile += 1
        for p in (p for p in punkte if p.get("Bereich") == bereich):
            seit = lies_datum(p.get("Seit", ""))
            werte = (p.get("Punkt", ""), p.get("Stand", ""), p.get("Wartet auf", ""),
                     seit.strftime("%d.%m.%Y") if seit else p.get("Seit", ""))
            for spalte, wert in zip(range(2, 6), werte):
                z = blatt.cell(row=zeile, column=spalte, value=wert)
                z.font = Font(name=SCHRIFT, size=10, color="404040" if spalte != 3 else DUNKEL,
                              bold=spalte == 3)
                z.alignment = Alignment(vertical="center", wrap_text=True,
                                        horizontal="center" if spalte in (3, 5) else "left")
                z.border = linie
            farbe = KURZ_FARBEN.get(p.get("Stand", ""))
            if farbe:
                blatt.cell(row=zeile, column=3).fill = PatternFill("solid", fgColor=farbe)
            blatt.row_dimensions[zeile].height = 20
            zeile += 1
        zeile += 1

    zelle(blatt, f"B{zeile}",
          "Bereit = erledigt und in Betrieb · Läuft = in Arbeit · Wartet = braucht eine andere Person · "
          "Offen = noch nicht begonnen oder Entscheid von Ingo. Details und Nachweise: Wochen_Todo.xlsx.",
          font=Font(name=SCHRIFT, size=9, italic=True, color=GRAU),
          alignment=Alignment(wrap_text=True, vertical="top"))
    blatt.merge_cells(f"B{zeile}:E{zeile}")
    blatt.row_dimensions[zeile].height = 26

    blatt.sheet_view.showGridLines = False
    blatt.freeze_panes = "A10"
    blatt.page_setup.orientation = "portrait"
    blatt.page_setup.fitToWidth = 1
    blatt.page_setup.fitToHeight = 0
    blatt.sheet_properties.pageSetUpPr.fitToPage = True
    blatt.print_title_rows = "9:9"
    mappe.save(KURZ_XLSX)


def meilenstein(wert: str) -> tuple[str, str]:
    """Symbol und Schriftfarbe fuer einen Entscheidungspunkt."""
    w = wert.strip()
    if w in ("", "-"):
        return "–  nicht nötig", GRAU
    if w.lower() == "offen":
        return "○  offen", "C00000"
    return f"✔  {w}", "2E8B57"


def rolle(wert: str) -> tuple[str, Font]:
    w = wert.strip()
    if w.lower().startswith("festlegen"):
        return "⚠ " + w, Font(name=SCHRIFT, size=10, color="C00000", italic=True)
    if w.lower().startswith("vorschlag"):
        return w, Font(name=SCHRIFT, size=10, color="2F5597", italic=True)
    return w, Font(name=SCHRIFT, size=10)


def kachel(blatt, spalte: int, zeile: int, titel: str, zahl, farbe: str, hinter: str,
           breite: int = 2) -> None:
    a, b = get_column_letter(spalte), get_column_letter(spalte + breite - 1)
    if breite > 1:
        blatt.merge_cells(f"{a}{zeile}:{b}{zeile}")
        blatt.merge_cells(f"{a}{zeile + 1}:{b}{zeile + 1}")
    fuell = PatternFill("solid", fgColor=hinter)
    for r in (zeile, zeile + 1):
        for c in range(spalte, spalte + breite):
            blatt.cell(row=r, column=c).fill = fuell
            blatt.cell(row=r, column=c).border = RAHMEN
    zelle(blatt, f"{a}{zeile}", titel, font=Font(name=SCHRIFT, size=9, color=GRAU, bold=True),
          alignment=Alignment(horizontal="center", vertical="bottom"))
    zelle(blatt, f"{a}{zeile + 1}", zahl, font=Font(name=SCHRIFT, size=22, bold=True, color=farbe),
          alignment=Alignment(horizontal="center", vertical="center"))


def schreibe_hermes(mappe, vorhaben: list[dict[str, str]], aufgaben: list[dict[str, str]]) -> None:
    blatt = neues_blatt(mappe, BLATT_HERMES, 0)
    blatt.sheet_view.showGridLines = False
    blatt.sheet_properties.tabColor = DUNKEL

    breiten = [8, 36, 13, 11, 22, 26, 13, 15, 13, 10, 46, 38, 42]
    for i, b in enumerate(breiten, start=1):
        blatt.column_dimensions[get_column_letter(i)].width = b
    letzte = get_column_letter(len(breiten))

    # Titelband
    blatt.merge_cells(f"A1:{letzte}1")
    zelle(blatt, "A1", "Vorhaben nach HERMES  ·  Übersicht für die Geschäftsleitung",
          font=Font(name=SCHRIFT, size=18, bold=True, color="FFFFFF"),
          fill=PatternFill("solid", fgColor=DUNKEL),
          alignment=Alignment(vertical="center", indent=1))
    blatt.row_dimensions[1].height = 36
    blatt.merge_cells(f"A2:{letzte}2")
    zelle(blatt, "A2",
          f"Stand {date.today():%d.%m.%Y}  ·  eine Zeile je Vorhaben  ·  "
          "Rollen mit «Vorschlag» sind aus der Dokumentation abgeleitet und noch nicht bestätigt",
          font=Font(name=SCHRIFT, size=10, italic=True, color="FFFFFF"),
          fill=PatternFill("solid", fgColor="2F5597"),
          alignment=Alignment(vertical="center", indent=1))
    blatt.row_dimensions[2].height = 20

    # Offene Aufgaben je Vorhaben aus dem Wochen-Todo
    offen = {v["ID"]: 0 for v in vorhaben}
    for a in aufgaben:
        if a.get("Vorhaben") in offen and not a.get("Status", "").startswith("Erledigt"):
            offen[a["Vorhaben"]] += 1

    ampeln = [v["Ampel"] for v in vorhaben]
    freigaben_fehlen = sum(1 for v in vorhaben if v["Freigabe"].strip().lower() == "offen")
    rollen_fehlen = sum(
        1 for v in vorhaben
        for k in ("Auftraggeber", "Anwendervertreter")
        if v[k].strip().lower().startswith("festlegen"))

    # Kennzahlkacheln
    blatt.row_dimensions[4].height = 18
    blatt.row_dimensions[5].height = 38
    kacheln = [
        ("VORHABEN", len(vorhaben), DUNKEL, HELL),
        ("● IM PLAN", ampeln.count("Grün"), *AMPEL["Grün"]),
        ("● ACHTUNG", ampeln.count("Gelb"), *AMPEL["Gelb"]),
        ("● KRITISCH", ampeln.count("Rot"), *AMPEL["Rot"]),
        ("● NEU", ampeln.count("Neu"), *AMPEL["Neu"]),
        ("FREIGABE FEHLT", freigaben_fehlen, "C00000", "F2F2F2"),
        ("ROLLE OFFEN", rollen_fehlen, "C55A11", "F2F2F2"),
    ]
    spalte = 1
    for titel, zahl, farbe, hinter in kacheln:
        breite = min(2, len(breiten) - spalte + 1)
        kachel(blatt, spalte, 4, titel, zahl, farbe, hinter, breite)
        spalte += breite

    # Tabelle
    kopf = ["ID", "Vorhaben", "HERMES-Stufe", "Ampel", "Auftraggeber", "Anwendervertreter",
            "① Freigabe", "② Abnahme", "③ Abschluss", "Offene Aufgaben",
            "Stand in einem Satz", "Nächster Entscheid", "Erledigt, wenn"]
    k = 7
    for i, name in enumerate(kopf, start=1):
        z = blatt.cell(row=k, column=i, value=name)
        z.font = Font(name=SCHRIFT, size=10, bold=True, color="FFFFFF")
        z.fill = PatternFill("solid", fgColor=DUNKEL)
        z.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        z.border = RAHMEN
    blatt.row_dimensions[k].height = 30

    for n, v in enumerate(vorhaben, start=1):
        r = k + n
        streifen = PatternFill("solid", fgColor="F7F9FC" if n % 2 else "FFFFFF")
        werte = [v["ID"], v["Vorhaben"], v["HERMES-Stufe"], None, None, None, None, None, None,
                 offen[v["ID"]], v["Stand in einem Satz"], v["Nächster Entscheid"], v["Erledigt, wenn"]]
        for c, w in enumerate(werte, start=1):
            z = blatt.cell(row=r, column=c, value=w)
            z.font = Font(name=SCHRIFT, size=10)
            z.fill = streifen
            z.border = RAHMEN
            z.alignment = Alignment(vertical="center", wrap_text=True)
        blatt.cell(row=r, column=1).font = Font(name=SCHRIFT, size=10, bold=True, color=DUNKEL)
        blatt.cell(row=r, column=2).font = Font(name=SCHRIFT, size=10, bold=True)

        stufe = blatt.cell(row=r, column=3)
        stufe.fill = PatternFill("solid", fgColor=STUFE_FARBEN.get(v["HERMES-Stufe"], "FFFFFF"))
        stufe.alignment = Alignment(horizontal="center", vertical="center")

        farbe, hinter = AMPEL.get(v["Ampel"], (GRAU, "FFFFFF"))
        amp = blatt.cell(row=r, column=4, value=f"●  {v['Ampel']}")
        amp.font = Font(name=SCHRIFT, size=11, bold=True, color=farbe)
        amp.fill = PatternFill("solid", fgColor=hinter)
        amp.alignment = Alignment(horizontal="center", vertical="center")

        for c, key in ((5, "Auftraggeber"), (6, "Anwendervertreter")):
            text, schrift = rolle(v[key])
            z = blatt.cell(row=r, column=c, value=text)
            z.font = schrift

        for c, key in ((7, "Freigabe"), (8, "Abnahme"), (9, "Abschluss")):
            text, farbe = meilenstein(v[key])
            z = blatt.cell(row=r, column=c, value=text)
            z.font = Font(name=SCHRIFT, size=10, bold=text.startswith("✔"), color=farbe)
            z.alignment = Alignment(horizontal="center", vertical="center")

        blatt.cell(row=r, column=10).alignment = Alignment(horizontal="center", vertical="center")
        # Hoehe aus der laengsten Textspalte schaetzen, sonst laeuft der Text ueber
        zeilen_max = max(
            math.ceil(len(str(blatt.cell(row=r, column=c).value or "")) / (breiten[c - 1] * 1.0))
            for c in (2, 11, 12, 13))
        blatt.row_dimensions[r].height = max(34, 15 * zeilen_max + 8)

    ende = k + len(vorhaben)
    blatt.conditional_formatting.add(
        f"J{k + 1}:J{ende}",
        DataBarRule(start_type="num", start_value=0, end_type="max", color="5B9BD5", showValue=True))
    blatt.freeze_panes = f"C{k + 1}"

    # Hilfsdaten fuer das Ampel-Diagramm, ausgeblendet rechts ausserhalb des Druckbereichs
    hs = len(breiten) + 3
    blatt.cell(row=k, column=hs, value="Ampel")
    blatt.cell(row=k, column=hs + 1, value="Anzahl")
    for i, name in enumerate(AMPEL_REIHE, start=1):
        blatt.cell(row=k + i, column=hs, value=name)
        blatt.cell(row=k + i, column=hs + 1, value=ampeln.count(name))
    for c in (hs, hs + 1):
        blatt.column_dimensions[get_column_letter(c)].hidden = True

    kuchen = PieChart()
    kuchen.title = "Ampel der Vorhaben"
    kuchen.add_data(Reference(blatt, min_col=hs + 1, min_row=k, max_row=k + len(AMPEL_REIHE)), titles_from_data=True)
    kuchen.set_categories(Reference(blatt, min_col=hs, min_row=k + 1, max_row=k + len(AMPEL_REIHE)))
    for i, name in enumerate(AMPEL_REIHE):
        punkt = DataPoint(idx=i)
        punkt.graphicalProperties.solidFill = AMPEL[name][0]
        punkt.graphicalProperties.line.solidFill = "FFFFFF"
        kuchen.series[0].dPt.append(punkt)
    kuchen.dataLabels = DataLabelList()
    kuchen.dataLabels.showVal = True
    kuchen.dataLabels.showSerName = False
    kuchen.dataLabels.showCatName = False
    kuchen.dataLabels.showLegendKey = False
    # Die Hilfsdaten sind ausgeblendet; ohne diese Einstellung zeichnet Excel sie nicht.
    kuchen.visible_cells_only = False
    kuchen.height, kuchen.width = 7.5, 10
    diagramm_zeile = ende + 2
    blatt.add_chart(kuchen, f"A{diagramm_zeile}")

    balken = BarChart()
    balken.type = "bar"
    balken.title = "Offene Aufgaben je Vorhaben"
    balken.add_data(Reference(blatt, min_col=10, min_row=k, max_row=ende), titles_from_data=True)
    balken.set_categories(Reference(blatt, min_col=1, min_row=k + 1, max_row=ende))
    balken.series[0].graphicalProperties.solidFill = "5B9BD5"
    balken.legend = None
    balken.y_axis.majorGridlines = None
    balken.x_axis.scaling.orientation = "maxMin"
    balken.dataLabels = DataLabelList()
    balken.dataLabels.showVal = True
    balken.dataLabels.showSerName = False
    balken.dataLabels.showCatName = False
    balken.dataLabels.showLegendKey = False
    balken.x_axis.delete = False
    balken.y_axis.delete = True
    balken.height, balken.width = 7.5, 16
    blatt.add_chart(balken, f"E{diagramm_zeile}")

    # Einstufung und Legende rechts neben den Diagrammen
    r0 = diagramm_zeile
    blatt.merge_cells(f"K{r0}:M{r0}")
    zelle(blatt, f"K{r0}", "So wird eingestuft (HERMES zugeschnitten)",
          font=Font(name=SCHRIFT, size=11, bold=True, color="FFFFFF"),
          fill=PatternFill("solid", fgColor=DUNKEL), alignment=Alignment(indent=1))
    stufen = [
        ("Stufe", "Aufwand", "Vorgehen"),
        ("Kleinauftrag", "unter ca. 5 Personentage", "Ticket, keine HERMES-Dokumente"),
        ("Vorhaben", "ca. 5 bis 40 Personentage", "Auftrag, Abnahme, Statuszeile"),
        ("Projekt", "über ca. 40 Personentage", "HERMES zugeschnitten, mit Risikoliste, eventuell Ausschuss"),
    ]
    for i, (a, b, c) in enumerate(stufen, start=1):
        for sp, w in zip("KLM", (a, b, c)):
            z = zelle(blatt, f"{sp}{r0 + i}", w, border=RAHMEN,
                      alignment=Alignment(vertical="center", wrap_text=True),
                      font=Font(name=SCHRIFT, size=10, bold=(i == 1)))
            if i == 1:
                z.fill = PatternFill("solid", fgColor=HELL)
            elif sp == "K":
                z.fill = PatternFill("solid", fgColor=STUFE_FARBEN[a])
    legende = [
        ("① Freigabe", "Der Auftraggeber hat den Projektauftrag unterschrieben."),
        ("② Abnahme", "Der Anwendervertreter bestätigt fachlich, das Ergebnis ist produktiv."),
        ("③ Abschluss", "Übergeben in den Betrieb, mit kurzer Betriebsdokumentation."),
        ("✔ / ○", "Erledigt mit Datum / noch offen."),
        ("⚠ festlegen", "Die Rolle ist noch niemandem zugeordnet."),
        ("● Ampel", "Grün im Plan · Gelb wartet auf Entscheid oder Zulieferung · Rot blockiert oder Termin überschritten · Grau neu, noch ohne Stand."),
    ]
    for i, (a, b) in enumerate(legende, start=len(stufen) + 2):
        zelle(blatt, f"K{r0 + i}", a, font=Font(name=SCHRIFT, size=10, bold=True, color=DUNKEL),
              alignment=Alignment(vertical="top"))
        blatt.merge_cells(f"L{r0 + i}:M{r0 + i}")
        zelle(blatt, f"L{r0 + i}", b, font=Font(name=SCHRIFT, size=10),
              alignment=Alignment(vertical="top", wrap_text=True))
        blatt.row_dimensions[r0 + i].height = 40 if len(b) > 80 else 28

    # Druck: quer, auf Seitenbreite
    blatt.page_setup.orientation = "landscape"
    blatt.page_setup.paperSize = blatt.PAPERSIZE_A4
    blatt.page_setup.fitToWidth = 1
    blatt.page_setup.fitToHeight = 0
    blatt.sheet_properties.pageSetUpPr.fitToPage = True
    blatt.print_area = f"A1:{letzte}{r0 + len(stufen) + len(legende) + 2}"
    blatt.print_title_rows = f"{k}:{k}"


def formular(mappe, name: str, titel: str, untertitel: str, abschnitte, farbe: str) -> None:
    """Einseitiges Formular. Ein Abschnitt ist eine Feldliste [(Feld, Hoehe)] oder eine
    Tabelle als Tupel (Kopfzeile, Anzahl Leerzeilen)."""
    blatt = neues_blatt(mappe, name)
    blatt.sheet_view.showGridLines = False
    blatt.sheet_properties.tabColor = farbe
    for sp, b in zip("ABCDE", (26, 22, 22, 22, 22)):
        blatt.column_dimensions[sp].width = b

    blatt.merge_cells("A1:E1")
    zelle(blatt, "A1", titel, font=Font(name=SCHRIFT, size=16, bold=True, color="FFFFFF"),
          fill=PatternFill("solid", fgColor=farbe), alignment=Alignment(vertical="center", indent=1))
    blatt.row_dimensions[1].height = 32
    blatt.merge_cells("A2:E2")
    zelle(blatt, "A2", untertitel, font=Font(name=SCHRIFT, size=9, italic=True, color=GRAU),
          alignment=Alignment(indent=1, wrap_text=True))
    blatt.row_dimensions[2].height = 28

    r = 4
    for ueberschrift, inhalt in abschnitte:
        blatt.merge_cells(f"A{r}:E{r}")
        zelle(blatt, f"A{r}", ueberschrift, font=Font(name=SCHRIFT, size=11, bold=True, color=farbe))
        for c in range(1, 6):
            blatt.cell(row=r, column=c).border = Border(bottom=Side(style="medium", color=farbe))
        r += 1
        if isinstance(inhalt, tuple):
            kopf, leer = inhalt
            for c, w in enumerate(kopf, start=1):
                zelle(blatt, f"{get_column_letter(c)}{r}", w, border=RAHMEN,
                      font=Font(name=SCHRIFT, size=10, bold=True),
                      fill=PatternFill("solid", fgColor=HELL),
                      alignment=Alignment(horizontal="center", wrap_text=True))
            r += 1
            for _ in range(leer):
                for c in range(1, len(kopf) + 1):
                    blatt.cell(row=r, column=c).border = RAHMEN
                blatt.row_dimensions[r].height = 22
                r += 1
        else:
            for feld, hoehe in inhalt:
                zelle(blatt, f"A{r}", feld, font=Font(name=SCHRIFT, size=10, bold=True),
                      alignment=Alignment(vertical="top", wrap_text=True),
                      fill=PatternFill("solid", fgColor="F2F2F2"), border=RAHMEN)
                blatt.merge_cells(f"B{r}:E{r}")
                for c in range(2, 6):
                    blatt.cell(row=r, column=c).border = RAHMEN
                blatt.cell(row=r, column=2).alignment = Alignment(vertical="top", wrap_text=True)
                blatt.row_dimensions[r].height = hoehe
                r += 1
        r += 1

    blatt.page_setup.orientation = "portrait"
    blatt.page_setup.paperSize = blatt.PAPERSIZE_A4
    blatt.page_setup.fitToWidth = 1
    blatt.page_setup.fitToHeight = 1
    blatt.sheet_properties.pageSetUpPr.fitToPage = True


def schreibe_vorlagen(mappe) -> None:
    formular(
        mappe, BLATT_AUFTRAG, "Projektauftrag  ·  eine Seite",
        "Für Vorhaben und Projekte nach HERMES (ab ca. 5 Personentage). Mit der Unterschrift des "
        "Auftraggebers ist der Entscheidungspunkt ① Freigabe erreicht.",
        [
            ("Vorhaben", [("Bezeichnung und ID", 22), ("HERMES-Stufe", 22),
                          ("Grober Aufwand (Personentage)", 22)]),
            ("Inhalt", [("Ausgangslage", 60), ("Ziel", 45), ("Nutzen (ein Satz genügt)", 30),
                        ("Nicht im Umfang", 30)]),
            ("Rollen", [("Auftraggeber (entscheidet, priorisiert)", 30),
                        ("Anwendervertreter (definiert fachlich, nimmt ab)", 30),
                        ("Projektleitung / Umsetzung", 24)]),
            ("Abnahmekriterien: Das Vorhaben ist fertig, wenn …",
             (["Nr.", "Kriterium", "", "Wie geprüft", ""], 4)),
            ("Freigabe", [("Termin Zielabnahme", 22), ("Datum, Unterschrift Auftraggeber", 34)]),
            ("Änderungsanträge nach der Freigabe (gehen über den Auftraggeber)",
             (["Datum", "Änderung", "Aufwand", "Entscheid", "Visum"], 3)),
        ],
        "2F5597",
    )
    formular(
        mappe, BLATT_ABNAHME, "Abnahmeprotokoll  ·  eine Seite",
        "Mit der Unterschrift des Anwendervertreters ist der Entscheidungspunkt ② Abnahme erreicht. "
        "Die Übergabe in den Betrieb schliesst das Vorhaben ab (③ Abschluss).",
        [
            ("Vorhaben", [("Bezeichnung und ID", 22), ("Abnahmegegenstand und Version", 30),
                          ("Datum der Abnahme", 22)]),
            ("Prüfung der Abnahmekriterien aus dem Projektauftrag",
             (["Nr.", "Kriterium", "✔ erfüllt / ✘ nicht", "Bemerkung", ""], 5)),
            ("Offene Mängel", (["Nr.", "Mangel", "Schwere", "Erledigen bis", "Wer"], 3)),
            ("Entscheid", [("☐ abgenommen\n☐ mit Auflagen\n☐ nicht abgenommen", 48), ("Auflagen", 40)]),
            ("Übergabe in den Betrieb", [("Betriebsdokumentation vorhanden (Ablage)", 30),
                                         ("Zuständig im Betrieb", 22)]),
            ("Unterschriften", [("Anwendervertreter, Datum", 34),
                                ("Auftraggeber (zur Kenntnis), Datum", 34)]),
        ],
        "2E8B57",
    )


def main() -> int:
    zeilen = lies_tsv(TSV)
    if not zeilen:
        print("TSV ist leer, nichts geschrieben.")
        return 1

    mappe = load_workbook(XLSX)
    # Das Datenblatt wird ueber den Namen gesucht, nicht ueber die Position.
    # Das Blatt `Zusammenfassung` wird von Hand gepflegt; es versehentlich zu
    # ueberschreiben war am 2026-09-03 ein Fehler.
    if BLATT_TODO not in mappe.sheetnames:
        print(f"Blatt '{BLATT_TODO}' nicht gefunden. Vorhanden: " + ", ".join(mappe.sheetnames))
        return 2

    eigene = {BLATT_TODO, BLATT_HERMES, BLATT_AUFTRAG, BLATT_ABNAHME}
    andere = [n for n in mappe.sheetnames if n not in eigene]

    schreibe_todo(mappe, zeilen)

    if VORHABEN_TSV.exists():
        vorhaben = als_dicts(lies_tsv(VORHABEN_TSV))
        aufgaben = als_dicts(zeilen)
        bekannte = {v["ID"] for v in vorhaben}
        fremd = sorted({a.get("Vorhaben", "") for a in aufgaben} - bekannte)
        if fremd:
            print("WARNUNG: Wochen-Todo verweist auf unbekannte Vorhaben: " + ", ".join(fremd))
        schreibe_hermes(mappe, vorhaben, aufgaben)
        schreibe_vorlagen(mappe)
        mappe.active = 0
        for blatt in mappe.worksheets:
            blatt.sheet_view.tabSelected = blatt.title == BLATT_HERMES
        print(f"{len(vorhaben)} Vorhaben in Blatt '{BLATT_HERMES}' geschrieben, Vorlagen erneuert.")



    mappe.save(XLSX)
    print(f"{len(zeilen) - 1} Zeilen in Blatt '{BLATT_TODO}' geschrieben.")

    if KURZ_TSV.exists():
        punkte = als_dicts(lies_tsv(KURZ_TSV))
        schreibe_status_kurz(punkte, als_dicts(zeilen))
        print(f"{len(punkte)} Punkte in {KURZ_XLSX.name} geschrieben.")
    if andere:
        print("Unveraendert erhalten: " + ", ".join(andere))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
